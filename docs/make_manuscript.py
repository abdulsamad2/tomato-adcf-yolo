"""Generate the manuscript .docx (python-docx). All numbers come from the project's frozen results files."""
import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"C:\Users\absamad\cs\tomato-adcf-yolo\docs\manuscript.docx")
FIG = Path(r"C:\Users\absamad\cs\posthoc")
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.3))
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 14), (2, 12), (3, 11)):
    h = doc.styles["Heading %d" % lvl]; h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = RGBColor(0, 0, 0); h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8); h.paragraph_format.space_after = Pt(4)

def runs(p, text):
    """**bold** and *italic* inline markup."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part: continue
        if part.startswith("**"): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*"): r = p.add_run(part[1:-1]); r.italic = True
        else: p.add_run(part)
    return p

def P(text, align=None, size=None, after=None):
    p = runs(doc.add_paragraph(), text)
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == "justify": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if size:
        for r in p.runs: r.font.size = Pt(size)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def H(text, lvl=1):
    return doc.add_heading(text, level=lvl)

def BL(items, style="List Bullet"):
    for it in items:
        p = runs(doc.add_paragraph(style=style), it); p.paragraph_format.space_after = Pt(2)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor); tcPr.append(s)

TCOUNT = [0]
def T(caption, header, rows, widths=None, note=None, size=8.5):
    TCOUNT[0] += 1
    cp = P("**Table %d.** %s" % (TCOUNT[0], caption), size=9.5, after=3)
    cp.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""; runs(c.paragraphs[0], "**%s**" % h); shade(c, "E7ECF2")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""; runs(c.paragraphs[0], str(v))
    for row in t.rows:
        for j, c in enumerate(row.cells):
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs: r.font.size = Pt(size)
            if widths: c.width = Cm(widths[j])
    if note: P(note, size=8.5, after=8)
    else: doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

FCOUNT = [0]
def F(path, caption, width=16.0):
    FCOUNT[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path), width=Cm(width))
    p.paragraph_format.keep_with_next = True
    P("**Figure %d.** %s" % (FCOUNT[0], caption), size=9.5, after=10)

# =====================================================================================================
P("**Beyond Metric Inflation: Source-Family Exposure Distorts Compression, Pretraining and Distillation "
  "Conclusions in Tomato Disease Detection**", align="center", size=15, after=10)
P("[Author names] · [Affiliations] · Corresponding author: [e-mail]", align="center", size=10, after=4)
P("*Draft manuscript, generated 13 October 2026. Bracketed items must be completed by the authors.*", align="center", size=9, after=14)

H("Abstract", 1)
P("Image-based plant-disease benchmarks are often distributed with offline augmented copies of each source "
  "photograph, and random file-level splits then place siblings of the same photograph in training and evaluation "
  "partitions. Prior work treats this mainly as score inflation. We ask whether it also changes the engineering "
  "conclusions drawn from a benchmark. Using the Tomato-Village object-detection data (14,368 files from 1,796 "
  "source photographs; every official validation file has a training sibling), we built a source-family-aware split "
  "and two development evaluation arms: a source-exposed arm and a source-independent arm. With five paired seeds "
  "and pre-registered decision rules we measured the effect of uniform width compression of YOLO11n, a "
  "localization-motivated width reallocation, COCO pretraining and localization-focused knowledge distillation under "
  "both arms. Source exposure magnified every estimated effect: shrinking to 0.50× width cost −12.2 AP under "
  "exposure but −1.8 AP under independent evaluation, and pretraining added +15.1 versus +3.9 AP. Under the same "
  "frozen selection rule, source-exposed validation rejected all reduced-width models, whereas source-independent "
  "validation selected a 1.57 M-parameter model (0.75× width) as equivalent to the 2.59 M-parameter model within "
  "±1.5 AP. Error decomposition showed that the dominant failure is lesions that receive no nearby detection at all "
  "(about half of ground-truth boxes), not disease misclassification (≤1.8 errors per 100 boxes); compression "
  "increased fully missed lesions and hurt small lesions most in relative terms. Distillation from a pretrained "
  "teacher reduced fully missed lesions and tightened boxes, raising the 0.75× model by +0.95 AP (95% CI +0.21 to "
  "+1.69) on the independent arm. On a locked test set evaluated once after all selections were frozen, the "
  "distilled 1.57 M model reached 16.08 AP, similar to the 2.59 M scratch model (16.19 AP), while the pretrained "
  "full-width model remained best (17.27 AP). Source-family exposure therefore biases intervention estimates and can "
  "change the model one would deploy.", align="justify")
P("**Keywords:** data leakage; evaluation validity; object detection; plant disease; model compression; knowledge "
  "distillation; pre-registration; YOLO.", size=10)

# =====================================================================================================
H("1. Introduction", 1)
P("Field detection of foliar disease is a natural deployment target for compact one-stage detectors, and the "
  "literature on lightweight plant-disease detectors is large [ADD REFS]. Most of these studies are developed and "
  "compared on public benchmarks whose images were augmented offline before release: each source photograph is "
  "accompanied by rotated, flipped, cropped or photometrically altered copies, and the copies are then split at the "
  "file level. Siblings of one photograph can then sit in both the training and the evaluation partitions. The "
  "consequence usually discussed is that reported accuracy is inflated, which is a familiar form of data leakage "
  "[12, 13].", align="justify")
P("This paper asks a different question: does source exposure also bias the *comparisons* that guide model design? "
  "If exposure magnifies the apparent benefit of more capacity, pretraining or an architectural change, then "
  "conclusions about which model to build and deploy may differ from those that an independent evaluation would "
  "support. We study this on Tomato-Village [1], a field-condition tomato disease detection dataset in which every "
  "source photograph is distributed as eight files.", align="justify")
P("Our contributions are:", after=2)
BL(["**Leakage characterisation.** We show that every one of the 1,796 source photographs has exactly seven augmented "
    "siblings and that all 2,875 official validation files have a sibling in official training. We then build a "
    "family-aware split and paired source-exposed and source-independent development arms.",
    "**Decision distortion.** Across six interventions (three compression levels, a width reallocation, COCO "
    "pretraining and distillation), exposure magnified every estimated effect. Under one frozen non-inferiority rule, "
    "the exposed protocol keeps the 2.59 M model while the independent protocol selects a 1.57 M model.",
    "**Mechanism.** A TIDE-style error decomposition [4], an IoU sweep and a lesion-size analysis show that the "
    "dominant failure under independent evaluation is missed lesions, not misclassification, and that compression "
    "harms small lesions most in relative terms.",
    "**A compact model with localization-focused distillation.** Distilling box distributions and P3/P4 spatial "
    "attention from a pretrained teacher partly restores localization in the 0.75× model without changing its "
    "inference graph. On a locked test set it matches the full-width scratch model but not the pretrained one.",
    "**Evidentiary discipline.** Primary analyses were pre-registered with hashed scripts, the test set was "
    "evaluated once after model selection was locked, and post-hoc analyses were specified and hashed before "
    "computation."])

H("2. Related work", 1)
P("**Leakage and evaluation validity.** Leakage, in which information unavailable at deployment time influences "
  "training or evaluation, is a recognised source of irreproducible machine-learning results [12, 13]. In plant "
  "pathology, studies using PlantVillage [14, 16] reported very high accuracies that later work attributed partly to "
  "acquisition artefacts and near-duplicate images [ADD REFS]. Group-aware splitting is the standard remedy when "
  "observations share a source. Our contribution is to measure how much the resulting bias propagates into "
  "*treatment effects* and *model-selection decisions*, rather than only into absolute scores.", align="justify")
P("**Lightweight detectors for plant disease.** Recent agricultural detectors combine YOLO-family backbones with "
  "attention modules, lightweight convolutions, modified necks or new box losses, typically reporting gains of a few "
  "AP points on random file-level splits [ADD REFS]. We use the YOLO11 family [2] and vary only width, one channel "
  "allocation, initialisation and the training loss, so that each effect can be attributed.", align="justify")
P("**Distillation for detection.** Knowledge distillation [5] has been adapted to dense detectors through feature "
  "imitation and localization distillation of box distributions [6]. The latter builds on the discretised box "
  "representation of Generalized Focal Loss [7], which YOLO11 also uses (distribution focal loss). We combine "
  "box-distribution distillation with channel-free spatial attention transfer [8].", align="justify")
P("**Error analysis.** TIDE [4] separates detection errors into classification, localization, duplicate, background "
  "and missed categories. We use a count-based variant at a fixed operating point.", align="justify")

# =====================================================================================================
H("3. Materials and methods", 1)
H("3.1 Data and source-family structure", 2)
P("We use Variant-c (object detection) of Tomato-Village [1]. The released files carry YOLO-format boxes for eight "
  "classes, two of which (Early Blight, Healthy) have no boxes. Files are named after their source photograph, with "
  "suffix _augN for derived copies. Grouping by source gives **1,796 families, each with exactly eight files** (one "
  "original and seven copies; 14,368 files in total). The official split has 11,493 training and 2,875 validation "
  "files. 1,484 families span both partitions, and **100% of validation files have a training sibling**; there is no "
  "official test split.", align="justify")
P("Comparing each copy's labels with its original's (box-matching tolerance 0.01 in normalised coordinates) "
  "classifies the 12,572 copies as follows. Photometric changes are invisible to labels.", after=2)
BL(["32.7% label-identical (photometric only);", "28.3% crop, scale or shift;",
    "14.2% horizontal flip and 14.2% vertical flip;", "7.2% 90° rotation and 3.5% 180° rotation."])

H("3.2 Family-aware split and evaluation arms", 2)
P("We grouped files by source photograph and merged near-duplicate photographs (difference hash, Hamming distance "
  "≤ 4; 59 pairs, giving 1,740 clusters). We then split clusters 70/15/15, stratified by majority class. This gives "
  "1,260 training photographs (10,080 files including their copies), 268 validation photographs and 268 test "
  "photographs (3,205 boxes, 261 clusters). Copies are kept in training only. The split manifest is released with "
  "the code.", align="justify")
P("For development we re-used the official validation files as two disjoint evaluation arms (Table 1). Both arms "
  "exclude every test family. They are scored from the same inference pass, so they differ only in which images "
  "enter the metric.", align="justify")
T("Development evaluation arms built from the official validation files.",
  ["Arm", "Definition", "Families", "Files", "Boxes"],
  [["A: source-exposed", "official-validation files whose source family is in our training split", "1,047", "2,028", "23,225"],
   ["B: source-independent", "official-validation files whose source family is in our validation split", "213", "411", "4,683"]],
  widths=[3.2, 7.6, 1.8, 1.6, 1.6],
  note="Arm B is the primary development metric. Arm A reproduces how a random file-level split behaves.")

H("3.3 Models and training protocol", 2)
P("All models are Ultralytics YOLO11n [2] variants trained with one protocol (Table 2). The protocol (YOLO11n, "
  "512 px, a fixed 24,570-update budget with a fixed checkpoint grid) was set in earlier development phases using "
  "validation data only (Section 3.6).", align="justify")
T("Training and evaluation protocol shared by all models.", ["Setting", "Value"],
  [["Detector / input", "YOLO11n family, 512 × 512"],
   ["Optimiser", "SGD, lr0 = 0.01, final lr factor 0.01 (linear), momentum 0.937, weight decay 5×10⁻⁴"],
   ["Batch / budget", "16 (nominal 64) · 39 epochs × 630 steps = 24,570 updates · 933 warm-up updates · mosaic off for the last 5 epochs"],
   ["Other", "Early stopping disabled · mixed precision · deterministic seeding · default Ultralytics augmentation"],
   ["Checkpoints", "Fixed at 6,300 / 11,970 / 18,270 / 24,570 updates; no best-validation selection"],
   ["Seeds", "0–4 for every model; comparisons are paired by seed"],
   ["Inference", "conf 0.001, NMS IoU 0.7, max 300 detections, top 100 per image, no test-time augmentation"],
   ["Scorer", "pycocotools AP@[.50:.95] [3], maxDets 100, 8 categories"],
   ["Operating point", "conf ≥ 0.25; greedy one-to-one matching at IoU 0.50 (strict 0.75). Localization recall = matched GT / all GT (class-agnostic); conditional class correctness = class-correct share of matched boxes"]],
  widths=[3.4, 12.4])
P("**Width.** Uniform channel width at 1.00×, 0.75×, 0.50× and 0.35× of YOLO11n, all trained from random "
  "initialisation. Initialising only some widths from COCO weights would confound width with initialisation: only "
  "1.00× can load 98% of the pretrained parameters, against 8–23% for the narrower models.", align="justify")
P("**LPWA-50.** This localization-preserving width allocation keeps the YOLO11n topology and moves channels from "
  "the stride-32 (P5) path to the stride-8/16 (P3/P4) paths. A parameter audit showed that the class branch is under "
  "2% of parameters, and that the box branch is already floored at 64 channels at every width. So the reallocatable "
  "budget sits in the backbone and neck, not the head. The allocation was solved by a rule fixed in advance: largest "
  "P3/P4 multiplier subject to parameters ≤ 0.811 M (uniform 0.50×) and GFLOPs ≤ 2.0. This gave P3/P4 ×1.37 and P5 "
  "×0.69, i.e. 0.787 M parameters and 1.965 GFLOPs.", align="justify")
P("**Pretrained reference.** YOLO11n initialised from COCO weights and trained with the same protocol.", align="justify")
P("**Localization-focused distillation.** The teacher is the pretrained 1.00× model (seed 0, final checkpoint), "
  "frozen. The student is the 0.75× model with an unchanged architecture and random initialisation identical to its "
  "baseline. The loss is", align="justify", after=2)
P("L = L_box + L_cls + L_dfl + λ_LD · g_dfl · L_LD + λ_AT · L_AT,", align="center")
P("where:", after=2)
BL(["**L_LD** is the KL divergence between teacher and student softmax distributions over the 16 DFL bins of each "
    "box side, averaged over the student's task-aligned-assigner foreground anchors [11] (temperature 1);",
    "**L_AT** sums, over P3 and P4, the squared distance between L2-normalised channel-mean squared activation maps "
    "[8]. These maps are channel-free, so no adapter layers are needed;",
    "**g_dfl** = 1.5 is the existing DFL gain, λ_LD = 1.0, and λ_AT = 1.25. λ_AT was set by a rule fixed before "
    "training: equal magnitude of the two distillation terms over the first ten training batches."])
P("No class-logit distillation was used and no trainable parameters were added. A dry run confirmed two things: "
  "with both λ set to zero the training loss reproduces the ordinary trainer exactly over 20 batches, and the "
  "teacher receives no gradients.", align="justify")

H("3.4 Statistics", 2)
BL(["**Seed uncertainty:** paired t over five seeds (df = 4).",
    "**Evaluation-set uncertainty:** a source-family bootstrap with 1,000 resamples, in which all files of a family "
    "move together. It holds trained checkpoints fixed. The two sources are reported separately and never pooled.",
    "**Equivalence margin:** ±1.5 AP, set before Phase 15 from the median 5-seed paired standard deviation of "
    "arm-B AP differences across 45 earlier model pairs (0.84 AP). At ±1.0 AP equivalence would only be attainable "
    "for |Δ| < 0.20.",
    "**Non-inferiority:** one-sided, Holm-corrected across the three reduced widths. A width is non-inferior when "
    "its Holm-rank-level lower confidence bound exceeds −1.5 AP.",
    "**Exposure distortion** of an intervention is D = Δ_A − Δ_B, where Δ is the paired effect on six-class AP. "
    "We emphasise absolute D and give ratios |Δ_A|/|Δ_B| only descriptively, because they are unstable when Δ_B is "
    "near zero."])

H("3.5 Evidentiary categories", 2)
P("Results fall into four categories, labelled throughout (Table 3).", align="justify")
T("Evidentiary status of each analysis.", ["Category", "Analyses", "Safeguard"],
  [["Prospective, pre-registered", "Width non-inferiority (Phase 15); LPWA-50 gates (Phase 16A); distillation gate (Phase 17)",
    "Rules, margins and script SHA-256 hashes frozen before training"],
   ["Locked test", "One evaluation of four models × five seeds", "Checkpoints and settings hashed; run lock; no gate and no selection"],
   ["Post-hoc, plan hashed", "Error decomposition, IoU sweep, operating curves, calibration, protocol decision comparison",
    "Written and hashed after the primary experiments but before computation; development arms only"],
   ["Exploratory", "Pretraining contrast; size thresholds; per-class effects; copy-type breakdown",
    "Labelled exploratory; pretraining-analysis plan written before scoring"]],
  widths=[3.6, 7.0, 5.2])

H("3.6 Development history and test-set disclosure", 2)
P("This study grew out of an attempt to improve YOLO necks with a gated detail–context fusion module. Under the "
  "family-aware split that module gave no benefit:", align="justify", after=2)
BL(["YOLO11s −0.4 AP (p = 0.75);", "5-fold grouped cross-validation 19.9 vs 20.1 AP;",
    "its learned gate stayed at α ≈ 0.50 everywhere."])
P("Subsequent development phases, all evaluated on validation data, chose YOLO11n over larger variants (no "
  "meaningful gain from size), 512 px input (equivalent to 640 px at ~64% of the GFLOPs) and the full three-scale "
  "head.", align="justify")
P("The original pipeline from that early phase automatically scored the family-aware test split for those earlier "
  "models, and one further test evaluation was made of 15 YOLO11s checkpoints. **None of the models reported here, "
  "and no protocol choice from the YOLO11n phases onward, was scored on or selected with the test split.** We "
  "therefore describe the test set as *held out from all models and decisions reported in this paper*, not as never "
  "opened.", align="justify")

# =====================================================================================================
H("4. Results", 1)
H("4.1 Width compression under independent evaluation", 2)
T("Width compression (arm B, 24,570 updates, mean ± SD over 5 seeds). Δ is versus 1.00×; LCB is the Holm-adjusted "
  "one-sided lower confidence bound.",
  ["Width", "Params", "GFLOPs", "AP arm B", "Δ", "LCB", "Verdict (margin −1.5)", "AP arm A"],
  [["1.00×", "2.59 M", "4.15", "13.65 ± 0.39", "—", "—", "reference", "29.76"],
   ["0.75×", "1.57 M", "2.71", "13.34 ± 0.54", "−0.31", "−1.32", "**non-inferior** (TOST-equivalent)", "23.96"],
   ["0.50×", "0.81 M", "1.55", "11.85 ± 0.52", "−1.80", "−2.80", "not established", "17.54"],
   ["0.35×", "0.54 M", "1.16", "10.45 ± 0.89", "−3.20", "−3.95", "not established", "13.21"]],
  widths=[1.4, 1.5, 1.4, 2.3, 1.3, 1.3, 4.2, 1.6])
P("Only 0.75× met the pre-registered criterion (Table 4). It was also equivalent to 1.00× under two-sided TOST "
  "(Holm p = 0.030). Batch-16 GPU latency was flat across widths (6.6–7.1 ms); batch-1 CPU latency fell from 27 ms "
  "to 17 ms. The narrower models were still improving between 18,270 and 24,570 updates, so their results are "
  "fixed-budget results.", align="justify")

H("4.2 The deployment decision depends on the evaluation protocol", 2)
P("Applying the identical frozen rule with arm-A instead of arm-B AP changes the selected model (Table 5, Figure 1). "
  "**Under the same frozen model-selection rule, source-exposed validation rejected all reduced-width models, whereas "
  "source-independent validation selected the 1.57 M-parameter model.**", align="justify")
T("Model selected by the Phase 15 rule under each validation protocol.",
  ["Validation", "0.75× (Δ, LCB)", "0.50×", "0.35×", "Selected"],
  [["Independent (arm B)", "−0.31, −1.32 ✓", "−1.80, −2.80 ✗", "−3.20, −3.95 ✗", "**0.75× (1.57 M)**"],
   ["Exposed (arm A)", "−5.79, −7.95 ✗", "−12.22, −14.06 ✗", "−16.54, −18.13 ✗", "**1.00× (2.59 M)**"]],
  widths=[3.2, 3.2, 3.0, 3.0, 3.4])
F(FIG / "fig_pareto.png", "Accuracy versus parameters (left) and GFLOPs (right). Circles and solid line: "
  "source-independent arm B. Squares and dashed line: source-exposed arm A. Uniform widths are joined; LPWA-50, "
  "distilled 0.75× and pretrained 1.00× are single markers. Exposure steepens the apparent cost of compression.")

H("4.3 Exposure distorts every intervention", 2)
T("Effect of each intervention on six-class AP under each arm (paired over 5 seeds, 95% CI).",
  ["Intervention", "Δ independent (B)", "Δ exposed (A)", "Distortion D = Δ_A − Δ_B", "|Δ_A| / |Δ_B|"],
  [["Width 0.75× (vs 1.00×)", "−0.31 [−1.19, 0.57]", "−5.79", "−5.48 [−6.58, −4.39]", "18.9"],
   ["Width 0.50×", "−1.80 [−2.80, −0.80]", "−12.22", "−10.42 [−12.56, −8.27]", "6.8"],
   ["Width 0.35×", "−3.20 [−4.17, −2.24]", "−16.54", "−13.34 [−15.03, −11.66]", "5.2"],
   ["LPWA-50 (vs uniform 0.50×)", "+0.26 [−0.51, 1.03]", "+1.58", "+1.32 [−1.08, 3.73]", "(6.2)"],
   ["COCO pretraining (vs scratch 1.00×)", "+3.93 [3.25, 4.61]", "+15.07", "+11.13 [9.60, 12.66]", "3.8"],
   ["Distillation (vs ordinary 0.75×)", "+0.95 [0.21, 1.69]", "+2.81", "+1.86 [−0.60, 4.32]", "3.0"]],
  widths=[4.6, 3.4, 2.2, 3.9, 1.9],
  note="Ratios are descriptive; the LPWA ratio is unstable because Δ_B is near zero.")
P("Source exposure consistently magnified apparent intervention effects (Table 6, Figure 2). The largest absolute "
  "distortions occurred for aggressive width reduction (−10.4 and −13.3 AP) and pretraining (+11.1 AP), where the "
  "confidence intervals exclude zero. For LPWA-50 and distillation, the direction was the same but the distortion "
  "intervals include zero. The pretraining contrast is exploratory: its analysis plan was written before scoring, "
  "but it was not part of a pre-registered phase.", align="justify")
F(FIG / "fig_distortion.png", "The same intervention measured under the source-exposed (red) and source-independent "
  "(blue) protocol. Effects are paired seed means of six-class AP.", width=14.5)

H("4.4 A localization-motivated reallocation does not pass its gate", 2)
P("LPWA-50 failed its pre-registered AP gate. Against uniform 0.50× it gave ΔAP_B = +0.26 [−0.51, +1.03], below "
  "the required +0.75. It did raise localization recall at IoU 0.50 by +1.31 [+0.33, +2.28], positive in all five "
  "seeds. Its gain under exposure was six times larger (+1.58), so a source-exposed study would have reported it as "
  "an architectural improvement. Architecture work stopped here under the frozen rule.", align="justify")

H("4.5 Localization-focused distillation", 2)
T("Distillation into the 0.75× model (arm B, 24,570 updates).",
  ["Model", "Params", "AP arm B", "AP arm A", "Loc. recall @0.50", "AP75"],
  [["Pretrained 1.00× (teacher family)", "2.59 M", "17.59", "44.82", "43.7", "15.22"],
   ["**Distilled 0.75×**", "1.57 M", "**14.30 ± 0.59**", "26.78", "36.3", "11.88"],
   ["From-scratch 1.00×", "2.59 M", "13.65 ± 0.39", "29.76", "36.9", "10.41"],
   ["Ordinary 0.75×", "1.57 M", "13.34 ± 0.54", "23.96", "34.6", "10.17"]],
  widths=[4.8, 1.8, 2.6, 2.0, 2.6, 1.8])
P("Distillation passed its gate (Table 7):", after=2)
BL(["**AP:** ΔAP_B = +0.95 [+0.21, +1.69], positive in all five seeds (+0.96, +0.62, +0.47, +1.97, +0.74; Figure 3).",
    "**Localization:** recall at IoU 0.50 rose by +1.75 [+0.79, +2.72], and AP75 by +1.70 [+0.44, +2.97].",
    "**Classification:** conditional class correctness was unchanged (+0.05).",
    "**Model:** parameters, GFLOPs, state-dict layout and file size are identical to the ordinary 0.75×, and so is "
    "latency.",
    "**Caveats:** the family-bootstrap interval for ΔAP_B is wider (−0.09 to +1.72); the student recovers only 22% "
    "of the gap to the pretrained teacher family; and the effect is again larger under exposure (+2.81)."])
F(FIG / "fig_seeds.png", "Seed-level paired effects on arm B. Green: ordinary 0.75× → distilled 0.75×. Orange: "
  "uniform 0.50× → LPWA-50. Left to right: AP50-95, localization recall at IoU 0.50, AP75.")

H("4.6 Mechanism: missed lesions dominate, and small lesions are most fragile", 2)
P("**Error decomposition (post-hoc, plan hashed).** At the operating point, about half of ground-truth lesions "
  "receive no detection overlapping them at IoU ≥ 0.1 (Table 8). Classification errors are rare in every model.", align="justify")
BL(["**Compression** raised the fully missed share by +3.2, +9.2 and +12.7 points for 0.75×, 0.50× and 0.35×.",
    "**Distillation** lowered it by −2.7 points, recovering 84% of the 0.75× increase, with no change in "
    "classification errors.",
    "**Pretraining** lowered it by −9.8 points."])
T("TIDE-style error counts on arm B at conf ≥ 0.25 (events per 100 ground-truth boxes; missed types as % of GT).",
  ["Model", "TP", "Loc. err.", "Cls. err.", "Cls+loc", "Dup.", "Backgr.", "Near-missed %", "Fully missed %"],
  [["1.00× scratch", "35.6", "9.7", "1.5", "0.8", "2.2", "8.6", "13.1", "**51.2**"],
   ["0.75×", "33.5", "8.5", "1.3", "0.8", "2.3", "7.6", "12.1", "**54.5**"],
   ["0.50×", "29.0", "7.3", "1.0", "0.6", "2.3", "5.9", "10.6", "**60.4**"],
   ["0.35×", "26.5", "6.6", "0.8", "0.5", "2.3", "5.5", "9.5", "**64.0**"],
   ["LPWA-50", "30.3", "7.4", "1.0", "0.6", "2.2", "6.4", "10.9", "58.8"],
   ["0.75× distilled", "35.2", "8.9", "1.3", "1.0", "2.3", "8.4", "13.0", "**51.8**"],
   ["1.00× pretrained", "42.3", "12.5", "1.8", "1.4", "1.7", "12.8", "16.3", "**41.4**"]],
  widths=[2.9, 1.2, 1.5, 1.5, 1.4, 1.2, 1.5, 2.2, 2.4])
P("**IoU sweep.** Distillation's AP gain was largest at IoU 0.70–0.80 (+1.70 and +1.32, against +0.78 at 0.50; "
  "Figure 4). This is consistent with tighter boxes and not only more detections. Its recall gain declined smoothly "
  "from +1.75 at IoU 0.50 to +0.47 at 0.90.", align="justify")
F(FIG / "fig_iou.png", "Six-class AP (left) and localization recall at conf ≥ 0.25 (right) as a function of the IoU "
  "threshold on arm B.")
P("**Lesion size (exploratory thresholds).** We binned ground truth by tertiles of √area on arm B (cut points 42.7 "
  "and 77.1 px; Table 9, Figure 5).", align="justify", after=2)
BL(["**Compression hurt small lesions most in relative terms.** 0.75× lost 23% of small-lesion recall against 2% of "
    "large-lesion recall; 0.50× lost 53% against 8%. Absolute losses were largest for medium lesions.",
    "**Distillation recovered** 59% (small), 75% (medium) and more than 100% (large) of the 0.75× loss.",
    "**Small lesions remain the unresolved problem:** even the pretrained model found only 14% of them."])
T("Localization recall at IoU 0.50 by lesion size on arm B (tertiles of √area).",
  ["Model", "Small (< 43 px)", "Medium", "Large (≥ 77 px)"],
  [["1.00× scratch", "9.6", "37.5", "63.5"], ["0.75×", "7.4", "34.1", "62.0"], ["0.50×", "4.5", "27.1", "58.1"],
   ["0.35×", "3.5", "23.3", "54.6"], ["0.75× distilled", "8.7", "36.6", "63.5"], ["1.00× pretrained", "14.3", "47.1", "69.6"]],
  widths=[4.0, 3.4, 3.0, 3.4])
F(FIG / "fig_size.png", "Localization recall at IoU 0.50 by lesion-size tertile on arm B.", width=13.0)
P("**Per class (exploratory).** The recall loss from compression and the recall gain from distillation appeared in "
  "nearly every class. Three classes have fewer than 100 arm-B boxes (Late Blight 78, Potassium 39, Nitrogen 32). "
  "The two deficiency classes among them are nearly undetectable (AP < 4) for every model (Supplementary Table S1).", align="justify")

H("4.7 Operating points and calibration", 2)
P("At matched false-positive rates of 2, 3 and 4 per family (Figure 6):", after=2)
BL(["the distilled 0.75× (29.1 / 32.6 / 35.0% recall) matched the 1.00× scratch model (28.9 / 32.4 / 35.0%);",
    "it exceeded the ordinary 0.75× by about 1.2 points;",
    "the pretrained model was best (32.6 / 36.5 / 39.3%)."])
P("Detection calibration error (D-ECE, 10 bins) was 0.026–0.030 for every model and was not affected by compression "
  "or distillation (Supplementary Figure S1).", align="justify")
F(FIG / "fig_operating.png", "Recall–false-positive operating curves on arm B for confidence thresholds 0.05–0.95.", width=11.5)

H("4.8 Leakage advantage depends on sibling similarity", 2)
P("This analysis is exploratory, for the 1.00× scratch model. The exposure advantage in localization recall (arm A "
  "minus arm B, same file type) was largest for images closest in pixels to a training sibling:", after=2)
BL(["originals: +13.6 points;", "photometric copies: +15.6;", "flips: about +14.5;",
    "crop or rescale copies: smallest, +6.4."])
P("Exposure benefit therefore follows visual similarity rather than being a uniform shift.", align="justify")

H("4.9 Locked test evaluation", 2)
P("After every selection was frozen, the four models were evaluated once on the 268 test photographs, with five "
  "seeds each and identical settings (Table 10).", align="justify")
T("Locked test results (mean ± SD over 5 seeds; one evaluation).",
  ["Model", "Params", "AP", "AP50", "AP75", "AR100", "Loc. R@0.50", "Loc. R@0.75", "Cls. correct", "FP/family"],
  [["**Distilled 0.75×**", "1.57 M", "**16.08 ± 0.96**", "33.07", "13.71", "31.29", "39.64", "22.05", "96.5", "2.63"],
   ["Ordinary 0.75×", "1.57 M", "14.95 ± 0.98", "31.23", "12.35", "29.58", "37.91", "20.79", "96.9", "2.47"],
   ["From-scratch 1.00×", "2.59 M", "16.19 ± 0.69", "33.69", "13.46", "31.55", "40.27", "22.12", "96.6", "2.68"],
   ["Pretrained 1.00×", "2.59 M", "**17.27 ± 0.39**", "34.66", "15.04", "30.11", "46.42", "25.09", "96.5", "3.43"]],
  widths=[2.9, 1.3, 2.2, 1.2, 1.2, 1.3, 1.6, 1.6, 1.5, 1.4], size=8,
  note="Paired Δ (distilled − reference), seed CI / family-bootstrap CI: vs ordinary 0.75× +1.14 [−0.42, 2.70] / "
       "[−0.21, 2.45]; vs scratch 1.00× −0.11 [−1.75, 1.53] / [−1.53, 1.17]; vs pretrained 1.00× −1.19 "
       "[−2.64, 0.26] / [−2.35, −0.08]. Localization recall @0.50 vs ordinary 0.75×: +1.73 [+0.29, +3.17].")
P("The test reproduced the development ranking: pretrained 1.00× > distilled 0.75× ≈ scratch 1.00× > ordinary "
  "0.75×.", align="justify", after=2)
BL(["The distillation gain (+1.14 AP) had the same size as on arm B, but its test-only interval includes zero.",
    "The localization-recall gain (+1.73) replicated with an interval excluding zero.",
    "The distilled 1.57 M model reached a similar observed mean AP to the 2.59 M scratch model, while the pretrained "
    "full-width model remained strongest. Under a pre-registered addendum, no model change followed from these results."])

# =====================================================================================================
H("5. Discussion", 1)
P("**From metric inflation to decision distortion.** Leakage is usually framed as an accounting error that inflates "
  "headline numbers. Our results show a more consequential effect: exposure changes the *differences* between "
  "models. A practitioner tuning capacity on a file-level split would conclude that even modest narrowing is "
  "costly, and would keep a model 1.65× larger than an independent evaluation supports. The same bias makes "
  "pretraining look four times as valuable and turns a null architectural change into an apparent gain. These are "
  "precisely the comparisons that architecture papers report, so exposure-aware evaluation should be a precondition "
  "for claims about plant-disease detector design.", align="justify")
P("**What the detectors actually fail at.** Under independent evaluation the detectors rarely confuse diseases. "
  "Given a found lesion, the class is right about 97% of the time in every condition. They fail by not finding "
  "lesions, and especially small ones. This shifts priorities away from classification-oriented additions (for "
  "example, attention for discriminating disease classes) toward high-resolution localization. It also explains "
  "why distilling box distributions helped while channel reallocation helped little: the student needed better "
  "localization targets, not merely more P3/P4 channels.", align="justify")
P("**Why the exposure gradient matters.** The leakage advantage was largest for copies most similar in pixels to a "
  "training image and smallest for crops or rescales. This supports interpreting the exposure effect as "
  "memorisation of specific photographs rather than an artefact of the split procedure, and suggests that "
  "label-preserving offline augmentation is the most harmful kind to split at the file level.", align="justify")
P("**Practical recommendations.**", after=2)
BL(["Release benchmarks with source identifiers.", "Split by source family.",
    "Report both exposed and independent results whenever a public split is compared with prior work.",
    "Base design decisions only on independent results."])

H("6. Limitations", 1)
BL(["One dataset with a 268-photograph test set (261 clusters); effects under about 1 AP have wide intervals, and "
    "test-only distillation evidence is directional.",
    "Severe class imbalance; two classes have no boxes and three have fewer than 100 development boxes.",
    "A fixed 24,570-update budget; smaller and from-scratch models had not plateaued.",
    "Desktop GPU/CPU timing only; no edge-device measurements and no speed claim (GPU latency was flat across widths).",
    "The test split was scored for earlier models outside this study (Section 3.6).",
    "Size thresholds, per-class effects, copy-type analyses and the pretraining contrast are exploratory.",
    "The distillation attention weight was set by a pre-specified scale-matching rule rather than tuned, and only "
    "one distillation configuration was tested."])

H("7. Conclusion", 1)
P("Source-family exposure distorts the estimated effects of compression, pretraining, architecture changes and "
  "distillation, and can reverse the resulting deployment decision. Under source-independent evaluation, the "
  "dominant failure of compact tomato-disease detectors is missed lesions rather than misclassification. "
  "Localization-focused distillation partly repairs that failure, giving a 1.57 M-parameter model that matches a "
  "2.59 M-parameter scratch model on a locked test set. Small-lesion localization remains the main open problem.",
  align="justify")

H("Data and code availability", 2)
P("The split manifest, pre-registrations with script hashes, analysis scripts, per-checkpoint results and trained "
  "weights will be released at [repository URL]. Tomato-Village is available from its authors [1].", align="justify")
H("Author contributions, funding and competing interests", 2)
P("[To be completed by the authors.]")

H("References", 1)
REFS = ["Gehlot M, Saxena RK, Gandhi GC. \"Tomato-Village\": a dataset for end-to-end tomato disease detection in a real-world environment. Multimedia Systems. 2023. doi:10.1007/s00530-023-01158-y.",
        "Jocher G, Qiu J. Ultralytics YOLO11. 2024. https://github.com/ultralytics/ultralytics.",
        "Lin T-Y, Maire M, Belongie S, et al. Microsoft COCO: Common Objects in Context. ECCV. 2014.",
        "Bolya D, Foley S, Hays J, Hoffman J. TIDE: A General Toolbox for Identifying Object Detection Errors. ECCV. 2020.",
        "Hinton G, Vinyals O, Dean J. Distilling the Knowledge in a Neural Network. arXiv:1503.02531. 2015.",
        "Zheng Z, Ye R, Wang P, et al. Localization Distillation for Dense Object Detection. CVPR. 2022.",
        "Li X, Wang W, Wu L, et al. Generalized Focal Loss: Learning Qualified and Distributed Bounding Boxes for Dense Object Detection. NeurIPS. 2020.",
        "Zagoruyko S, Komodakis N. Paying More Attention to Attention: Improving the Performance of Convolutional Neural Networks via Attention Transfer. ICLR. 2017.",
        "Schuirmann DJ. A comparison of the two one-sided tests procedure and the power approach for assessing the equivalence of average bioavailability. J Pharmacokinet Biopharm. 1987;15:657–680.",
        "Holm S. A simple sequentially rejective multiple test procedure. Scand J Stat. 1979;6:65–70.",
        "Feng C, Zhong Y, Gao Y, Scott MR, Huang W. TOOD: Task-aligned One-stage Object Detection. ICCV. 2021.",
        "Kapoor S, Narayanan A. Leakage and the reproducibility crisis in machine-learning-based science. Patterns. 2023;4(9):100804.",
        "Kaufman S, Rosset S, Perlich C, Stitelman O. Leakage in data mining: formulation, detection, and avoidance. ACM Trans Knowl Discov Data. 2012;6(4):15.",
        "Hughes DP, Salathé M. An open access repository of images on plant health to enable the development of mobile disease diagnostics. arXiv:1511.08060. 2015.",
        "Efron B, Tibshirani RJ. An Introduction to the Bootstrap. Chapman & Hall; 1993.",
        "Mohanty SP, Hughes DP, Salathé M. Using deep learning for image-based plant disease detection. Front Plant Sci. 2016;7:1419."]
for k, r in enumerate(REFS, 1):
    p = P("[%d] %s" % (k, r), size=9.5, after=2)
P("*[ADD REFS] markers indicate where recent lightweight agricultural detectors and PlantVillage bias studies should be "
  "cited; verify every reference before submission.*", size=9)

doc.add_page_break()
H("Supplementary material", 1)
T("S1. Per-class results on arm B (exploratory; mean over 5 seeds). ⚑ = fewer than 100 ground-truth boxes.",
  ["Class", "GT boxes", "AP 1.00×", "AP 0.75×", "AP 0.75× distilled", "Δ recall, width 0.75×", "Δ recall, distillation"],
  [["Leaf Miner", "3,138", "18.0", "17.1", "17.9", "−2.8", "+1.9"],
   ["Spotted Wilt Virus", "899", "19.5", "19.4", "20.2", "−0.1", "+0.0"],
   ["Magnesium Deficiency", "497", "22.1", "20.9", "23.1", "−3.4", "+3.3"],
   ["Late Blight ⚑", "78", "16.9", "17.2", "21.1", "−0.8", "+3.8"],
   ["Potassium Deficiency ⚑", "39", "3.3", "3.7", "2.4", "−5.6", "+4.1"],
   ["Nitrogen Deficiency ⚑", "32", "2.1", "1.6", "1.1", "−3.1", "+5.0"]],
  widths=[3.8, 1.8, 1.7, 1.7, 2.4, 2.3, 2.3])
F(FIG / "fig_calibration.png", "(S1) Reliability diagram on arm B, seed 0, with detection-ECE per model.", width=10.0)
T("S2. Inferred augmentation type of the 12,572 released copies, and exploratory exposure advantage in localization "
  "recall at IoU 0.50 (1.00× scratch, 5 seeds).",
  ["File type", "Share of copies", "Recall arm A", "Recall arm B", "Exposure advantage"],
  [["Original photograph", "—", "53.9", "40.3", "+13.6"], ["Label-identical (photometric)", "32.7%", "50.8", "35.2", "+15.6"],
   ["Horizontal flip", "14.2%", "53.3", "38.8", "+14.5"], ["Vertical flip", "14.2%", "49.8", "35.3", "+14.6"],
   ["Crop / scale / shift", "28.3%", "45.9", "39.4", "+6.4"], ["90° rotation", "7.2%", "43.3–48.5", "30.6–33.4", "+12.6–15.1"],
   ["180° rotation", "3.5%", "42.0", "35.5", "+6.5"]],
  widths=[4.6, 2.6, 2.4, 2.4, 3.2])
P("Additional supplementary items available in the repository: full IoU-sweep tables, per-checkpoint learning "
  "curves, pre-registration documents (Phases 15, 16A, 17, locked test) and the post-hoc analysis plan with its "
  "SHA-256 hash.", size=10)

for z in doc.settings.element.findall(qn("w:zoom")):
    z.set(qn("w:percent"), "100")
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("saved", OUT, "tables", TCOUNT[0] + 0, "figures", FCOUNT[0])
